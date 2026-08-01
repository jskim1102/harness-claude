from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import random
from typing import Iterable

from openpyxl import load_workbook


@dataclass(frozen=True)
class BulkRecord:
    rcp_id: str
    bulk_type: str | None
    functional: str | None
    functional_uv: str | None
    preservative: str | None
    stability: str | None
    product_group: str | None
    product_midgroup: str | None
    product_subtype: str | None
    form_type: str | None
    viscosity: str | None
    spindle: str | None
    rpm: str | None
    cutting_hardness: str | None
    penetration_hardness: str | None
    melting_point: str | None
    ph: str | None
    leave_on: str | None
    rinse_off: str | None
    gas_filling: str | None


@dataclass(frozen=True)
class RecipeLine:
    rcp_id: str
    phase: str
    ingredient_code: str
    ingredient_name: str
    amount: float


def load_dataset(data_dir: str | Path) -> tuple[list[BulkRecord], list[RecipeLine]]:
    data_path = Path(data_dir)
    return (
        load_bulk_records(data_path / "bulk_makeup.xlsx"),
        load_recipe_lines(data_path / "reciepe_makeup.xlsx"),
    )


def load_anonymized_dataset(
    data_dir: str | Path,
    seed: int = 20260707,
) -> tuple[list[BulkRecord], list[RecipeLine]]:
    records, lines, _sample_labels = _load_anonymized_parts(data_dir, seed)
    return records, lines


def load_anonymized_web_dataset(
    data_dir: str | Path,
    seed: int = 20260707,
) -> tuple[list[BulkRecord], list[RecipeLine], dict[str, str]]:
    return _load_anonymized_parts(data_dir, seed)


def _load_anonymized_parts(
    data_dir: str | Path,
    seed: int,
) -> tuple[list[BulkRecord], list[RecipeLine], dict[str, str]]:
    data_path = Path(data_dir)
    bulk_path = data_path / "bulk_makeup.xlsx"
    recipe_path = data_path / "reciepe_makeup.xlsx"

    workbook = load_workbook(bulk_path, read_only=False, data_only=True)
    sheet = workbook.active
    headers = _flatten_bulk_headers(sheet)
    index = {name: pos for pos, name in enumerate(headers)}

    raw_rows: list[tuple[str, tuple[object, ...]]] = []
    seen: set[str] = set()
    for row in sheet.iter_rows(min_row=3, values_only=True):
        values = tuple(row)
        rcp_id = _get(values, index, "RCP_ID(벌크품번)")
        if not rcp_id:
            continue
        raw_key = str(rcp_id)
        if raw_key in seen:
            continue
        seen.add(raw_key)
        raw_rows.append((raw_key, values))

    anonymous_ids = [f"A{number:05d}" for number in range(len(raw_rows))]
    random.Random(seed).shuffle(anonymous_ids)
    raw_to_anon = {
        raw_key: anonymous_id
        for (raw_key, _row), anonymous_id in zip(raw_rows, anonymous_ids)
    }
    sample_labels = {
        raw_to_anon[raw_key]: f"{raw_key}|{_get(row, index, '제품명') or ''}"
        for raw_key, row in raw_rows
    }

    records = [
        _bulk_record_from_row(raw_to_anon[raw_key], row, index)
        for raw_key, row in raw_rows
    ]

    recipe_workbook = load_workbook(recipe_path, read_only=True, data_only=True)
    recipe_sheet = recipe_workbook.active
    rows = recipe_sheet.iter_rows(min_row=1, values_only=True)
    header = next(rows)
    recipe_index = {name: pos for pos, name in enumerate(header)}
    rcp_key = "RCP_ID" if "RCP_ID" in recipe_index else "RCP_ID(벌크품번)"

    lines: list[RecipeLine] = []
    for row in rows:
        padded = tuple(row) + (None,) * len(header)
        raw_rcp_id = padded[recipe_index[rcp_key]]
        code = padded[recipe_index["원료코드"]]
        if raw_rcp_id is None or code is None:
            continue
        anonymous_id = raw_to_anon.get(str(raw_rcp_id))
        if anonymous_id is None:
            continue
        lines.append(
            RecipeLine(
                rcp_id=anonymous_id,
                phase=_phase(padded[recipe_index["상"]]),
                ingredient_code=str(code),
                ingredient_name=str(padded[recipe_index["원료명"]] or ""),
                amount=float(padded[recipe_index["함량"]] or 0.0),
            )
        )
    return records, lines, sample_labels


def load_bulk_records(path: str | Path) -> list[BulkRecord]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook.active
    headers = _flatten_bulk_headers(sheet)
    index = {name: pos for pos, name in enumerate(headers)}
    records: list[BulkRecord] = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        rcp_id = _get(row, index, "RCP_ID(벌크품번)")
        if not rcp_id:
            continue
        records.append(
            BulkRecord(
                rcp_id=str(rcp_id),
                bulk_type=_optional(_get(row, index, "벌크유형")),
                functional=_optional(_get(row, index, "기능성_벌크품목정보등록 선택")),
                functional_uv=_optional(_get(row, index, "기능성_유기 / 무기 / 유무기 선택")),
                preservative=_optional(_get(row, index, "방부력")),
                stability=_optional(_get(row, index, "안정도")),
                product_group=_optional(_get(row, index, "제품유형_대그룹")),
                product_midgroup=_optional(_get(row, index, "제품유형_중그룹")),
                product_subtype=_optional(_get(row, index, "제품유형_소그룹")),
                form_type=_optional(_get(row, index, "제품유형_제형타입")),
                viscosity=_optional(_get(row, index, "숙성기준_점도")),
                spindle=_optional(_get(row, index, "숙성기준_Spindle #")),
                rpm=_optional(_get(row, index, "숙성기준_rpm")),
                cutting_hardness=_optional(_get(row, index, "숙성기준_절단경도")),
                penetration_hardness=_optional(_get(row, index, "숙성기준_침입경도(kg)")),
                melting_point=_optional(_get(row, index, "숙성기준_융점")),
                ph=_optional(_get(row, index, "pH")),
                leave_on=_optional(_get(row, index, "Leave-on")),
                rinse_off=_optional(_get(row, index, "Rinse-off")),
                gas_filling=_optional(_get(row, index, "가스충진")),
            )
        )
    return records


def _bulk_record_from_row(
    anonymous_id: str,
    row: Iterable[object],
    index: dict[str, int],
) -> BulkRecord:
    return BulkRecord(
        rcp_id=anonymous_id,
        bulk_type=_optional(_get(row, index, "벌크유형")),
        functional=_optional(_get(row, index, "기능성_벌크품목정보등록 선택")),
        functional_uv=_optional(_get(row, index, "기능성_유기 / 무기 / 유무기 선택")),
        preservative=_optional(_get(row, index, "방부력")),
        stability=_optional(_get(row, index, "안정도")),
        product_group=_optional(_get(row, index, "제품유형_대그룹")),
        product_midgroup=_optional(_get(row, index, "제품유형_중그룹")),
        product_subtype=_optional(_get(row, index, "제품유형_소그룹")),
        form_type=_optional(_get(row, index, "제품유형_제형타입")),
        viscosity=_optional(_get(row, index, "숙성기준_점도")),
        spindle=_optional(_get(row, index, "숙성기준_Spindle #")),
        rpm=_optional(_get(row, index, "숙성기준_rpm")),
        cutting_hardness=_optional(_get(row, index, "숙성기준_절단경도")),
        penetration_hardness=_optional(_get(row, index, "숙성기준_침입경도(kg)")),
        melting_point=_optional(_get(row, index, "숙성기준_융점")),
        ph=_optional(_get(row, index, "pH")),
        leave_on=_optional(_get(row, index, "Leave-on")),
        rinse_off=_optional(_get(row, index, "Rinse-off")),
        gas_filling=_optional(_get(row, index, "가스충진")),
    )


def load_recipe_lines(path: str | Path) -> list[RecipeLine]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(min_row=1, values_only=True)
    header = next(rows)
    index = {name: pos for pos, name in enumerate(header)}
    rcp_key = "RCP_ID" if "RCP_ID" in index else "RCP_ID(벌크품번)"
    lines: list[RecipeLine] = []
    for row in rows:
        padded = tuple(row) + (None,) * len(header)
        rcp_id = padded[index[rcp_key]]
        code = padded[index["원료코드"]]
        if rcp_id is None or code is None:
            continue
        lines.append(
            RecipeLine(
                rcp_id=str(rcp_id),
                phase=_phase(padded[index["상"]]),
                ingredient_code=str(code),
                ingredient_name=str(padded[index["원료명"]] or ""),
                amount=float(padded[index["함량"]] or 0.0),
            )
        )
    return lines


def _flatten_bulk_headers(sheet) -> list[str]:
    row1 = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    row2 = [sheet.cell(2, col).value for col in range(1, sheet.max_column + 1)]
    current = None
    filled_row1: list[object | None] = []
    for value in row1:
        if value is not None:
            current = value
        filled_row1.append(current)
    return [
        str(top) if bottom is None else f"{top}_{bottom}"
        for top, bottom in zip(filled_row1, row2)
    ]


def _get(row: Iterable[object], index: dict[str, int], key: str):
    values = tuple(row)
    pos = index[key]
    return values[pos] if pos < len(values) else None


def _optional(value: object | None) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _phase(value: object | None) -> str:
    import unicodedata

    return unicodedata.normalize("NFKC", str(value or "")).strip()
